#!/usr/bin/env python3
"""
Prueba offline del flujo de LOTE de facturas (sin red, sin Telegram real).

Simula updates de Telegram con un request falso y verifica:
  A. Lote por menú: las N fotos enviadas se ACUMULAN en el lote (no se
     reinicia la conversación por el entry point con allow_reentry).
  B. Álbum directo (media_group_id): se procesa como lote automático.
  C. Los textos con Markdown escapan los caracteres especiales del nombre
     del proveedor (evita BadRequest: Can't parse entities).

Uso:  python test_batch_flow.py
"""

import asyncio
import io
import itertools
import json
import os
import sys
import tempfile
import time

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Config de entorno ANTES de importar bot.py
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123456:TESTTOKEN")
os.environ.setdefault("ANTHROPIC_API_KEY", "test-key")
_tmpdb = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
_tmpdb.close()
os.environ["DB_PATH"] = _tmpdb.name

import bot  # noqa: E402

from telegram import Update  # noqa: E402
from telegram.ext import ExtBot  # noqa: E402
from telegram.request import BaseRequest  # noqa: E402

UID = 111  # chat privado: chat_id == user_id


# ──────────────────────────────────────────────────────────────
# Bot API falsa
# ──────────────────────────────────────────────────────────────

class FakeRequest(BaseRequest):
    """Responde localmente a los métodos de la Bot API que usa el flujo."""

    def __init__(self):
        self._next_id = itertools.count(1000)
        self.last_message_id = None
        self.calls = []  # (metodo, params)

    async def initialize(self):
        pass

    async def shutdown(self):
        pass

    def _message(self, params):
        mid = params.get("message_id") or next(self._next_id)
        self.last_message_id = mid
        return {
            "message_id": mid,
            "date": int(time.time()),
            "chat": {"id": int(params.get("chat_id", UID)),
                     "type": "private", "first_name": "Ricardo"},
            "from": {"id": 42, "is_bot": True, "first_name": "Gabi"},
            "text": params.get("text", ""),
        }

    async def do_request(self, url, method, request_data=None,
                         read_timeout=None, write_timeout=None,
                         connect_timeout=None, pool_timeout=None):
        # Descarga de archivos (retrieve)
        if "/file/bot" in url:
            return 200, b"FAKEIMGDATA"

        api = url.rsplit("/", 1)[-1]
        params = dict(request_data.parameters) if request_data else {}
        self.calls.append((api, params))

        if api == "getMe":
            result = {"id": 42, "is_bot": True, "first_name": "Gabi",
                      "username": "gabi_bot",
                      "can_join_groups": True,
                      "can_read_all_group_messages": False,
                      "supports_inline_queries": False}
        elif api in ("sendMessage", "editMessageText"):
            result = self._message(params)
        elif api == "getFile":
            result = {"file_id": params["file_id"],
                      "file_unique_id": "u" + str(params["file_id"]),
                      "file_path": "photos/fake.jpg"}
        elif api in ("answerCallbackQuery", "deleteMessage", "deleteWebhook"):
            result = True
        else:
            result = True
        return 200, json.dumps({"ok": True, "result": result}).encode()


# ──────────────────────────────────────────────────────────────
# Fabricar updates
# ──────────────────────────────────────────────────────────────

_upd_id = itertools.count(1)
_msg_id = itertools.count(1)


def _user():
    return {"id": UID, "is_bot": False, "first_name": "Ricardo"}


def _chat():
    return {"id": UID, "type": "private", "first_name": "Ricardo"}


def upd_text(text):
    return {"update_id": next(_upd_id),
            "message": {"message_id": next(_msg_id), "date": int(time.time()),
                        "chat": _chat(), "from": _user(), "text": text}}


def upd_photo(fid, media_group=None):
    msg = {"message_id": next(_msg_id), "date": int(time.time()),
           "chat": _chat(), "from": _user(),
           "photo": [{"file_id": fid, "file_unique_id": "u" + fid,
                      "width": 800, "height": 600}]}
    if media_group:
        msg["media_group_id"] = media_group
    return {"update_id": next(_upd_id), "message": msg}


def upd_pdf(fid, name="facturas.pdf"):
    return {"update_id": next(_upd_id),
            "message": {"message_id": next(_msg_id), "date": int(time.time()),
                        "chat": _chat(), "from": _user(),
                        "document": {"file_id": fid,
                                     "file_unique_id": "u" + fid,
                                     "file_name": name,
                                     "mime_type": "application/pdf"}}}


def upd_cq(data, message_id):
    return {"update_id": next(_upd_id),
            "callback_query": {
                "id": str(next(_upd_id)), "from": _user(),
                "chat_instance": "ci-1", "data": data,
                "message": {"message_id": message_id, "date": int(time.time()),
                            "chat": _chat(),
                            "from": {"id": 42, "is_bot": True,
                                     "first_name": "Gabi"},
                            "text": "x"}}}


# ──────────────────────────────────────────────────────────────
# Stub de extracción con IA (sin API)
# ──────────────────────────────────────────────────────────────

_ncf_n = itertools.count(1)


async def fake_extract(image_bytes, filename):
    n = next(_ncf_n)
    return {
        "_filename": filename, "_error": None,
        "nombre_proveedor": f"PROVEEDOR {n}",
        "rnc": "101000001", "ncf": f"B01000000{n:02d}",
        "fecha_comprobante": "2026-06-15",
        "monto_sin_itbis": 100.0, "itbis": 18.0, "propina": 0.0,
        "total_facturado": 118.0, "metodo_pago": "EFECTIVO",
        "nivel_confianza": "ALTO", "observaciones": "",
        "_qr_verified": False, "_warnings": [], "_needs_review": False,
    }


# ──────────────────────────────────────────────────────────────
# Escenarios
# ──────────────────────────────────────────────────────────────

FAILS = []


def check(cond, label):
    mark = "OK " if cond else "FALLO"
    print(f"  [{mark}] {label}")
    if not cond:
        FAILS.append(label)


async def run():
    bot.extract_invoice = fake_extract
    bot.init_db()

    fake = FakeRequest()
    tbot = ExtBot(token=os.environ["TELEGRAM_BOT_TOKEN"],
                  request=fake, get_updates_request=FakeRequest())
    app = bot.build_application(bot=tbot)
    conv = app.handlers[0][0]

    await app.initialize()

    async def drive(data):
        await app.process_update(Update.de_json(data, app.bot))

    def state():
        return conv._conversations.get((UID, UID))

    ud = app.user_data[UID]

    # Desbloquear con la contraseña
    await drive(upd_text(bot.BOT_PASSWORD))
    check(ud.get("unlocked") is True, "contraseña desbloquea el bot")

    print("\n— Escenario A: lote por menú, álbum de 3 fotos —")
    await drive(upd_text(bot.BTN_NUEVA))
    check(state() == bot.S_MODE, "menú de modo (S_MODE)")
    await drive(upd_cq("mode_batch", fake.last_message_id))
    check(state() == bot.S_LOCATION, "lote → pregunta ubicación")
    await drive(upd_cq(f"loc_{bot.LOCATIONS[0]}", fake.last_message_id))
    await drive(upd_cq(f"cat_{bot.CATEGORIES[0]}", fake.last_message_id))
    check(state() == bot.S_COLLECT, "categoría → recolectando (S_COLLECT)")

    for i in range(1, 4):
        await drive(upd_photo(f"A{i}", media_group="ALBUM1"))
    ids = ud.get("batch_photo_ids") or []
    check(len(ids) == 3,
          f"las 3 fotos del álbum quedan en el lote (hay {len(ids)})")
    check(state() == bot.S_COLLECT,
          f"sigue recolectando, no se reinició (estado={state()})")

    await drive(upd_cq("batch_done", fake.last_message_id))
    items = ud.get("batch_items") or []
    check(len(items) == 3, f"lote procesado con 3 facturas (hay {len(items)})")
    check(state() == bot.S_BATCH_LIST, "lista del lote (S_BATCH_LIST)")

    # Revisar la #2, aceptarla desde la tarjeta, aceptar el resto
    await drive(upd_cq("brev_1", fake.last_message_id))
    check(state() == bot.S_CONFIRM, "🔍 Revisar abre la tarjeta (S_CONFIRM)")
    await drive(upd_cq("bconf_accept", fake.last_message_id))
    check(state() == bot.S_BATCH_LIST, "aceptar desde tarjeta vuelve a la lista")
    await drive(upd_cq("baccall", fake.last_message_id))
    guardadas = bot.get_facturas("2026-06")
    check(len(guardadas) == 3,
          f"las 3 facturas quedaron guardadas en 2026-06 (hay {len(guardadas)})")
    check(state() is None, "lote terminado cierra la conversación")

    print("\n— Escenario B: álbum directo (sin menú) —")
    for i in range(1, 4):
        await drive(upd_photo(f"B{i}", media_group="ALBUM2"))
    ids = ud.get("batch_photo_ids") or []
    check(ud.get("mode") == "batch_photo", "álbum directo entra en modo lote")
    check(len(ids) == 3, f"las 3 fotos del álbum acumuladas (hay {len(ids)})")
    check(state() == bot.S_LOCATION, "pregunta ubicación una sola vez")
    await drive(upd_cq(f"loc_{bot.LOCATIONS[1]}", fake.last_message_id))
    await drive(upd_cq(f"cat_{bot.CATEGORIES[1]}", fake.last_message_id))
    check(state() == bot.S_COLLECT, "pasa a recolectar con las fotos ya contadas")
    await drive(upd_cq("batch_done", fake.last_message_id))
    items = ud.get("batch_items") or []
    check(len(items) == 3, f"lote del álbum procesado (hay {len(items)})")
    await drive(upd_cq("baccall", fake.last_message_id))
    guardadas = bot.get_facturas("2026-06")
    check(len(guardadas) == 6,
          f"total 6 facturas guardadas (hay {len(guardadas)})")

    print("\n— Escenario D: foto individual (sin álbum) —")
    await drive(upd_photo("SOLO1"))  # sin media_group_id
    check(ud.get("mode") == "single", "foto suelta entra en modo individual")
    check(state() == bot.S_LOCATION, "pregunta ubicación de la compra")
    await drive(upd_cq(f"loc_{bot.LOCATIONS[0]}", fake.last_message_id))
    await drive(upd_cq(f"cat_{bot.CATEGORIES[0]}", fake.last_message_id))
    check(state() == bot.S_CONFIRM, "muestra la tarjeta de revisión")
    await drive(upd_cq("confirm_accept", fake.last_message_id))
    guardadas = bot.get_facturas("2026-06")
    check(len(guardadas) == 7,
          f"la factura individual quedó guardada (hay {len(guardadas)})")
    await drive(upd_cq("fin_lote", fake.last_message_id))
    check(state() is None, "Terminar cierra la conversación")

    print("\n— Escenario E: PDF directo (1 factura por página) —")
    bot.render_pdf_pages = lambda pdf_bytes, **kw: [b"PG1", b"PG2"]
    await drive(upd_pdf("PDF1"))
    check(ud.get("mode") == "batch_pdf", "PDF directo entra en modo lote PDF")
    check(state() == bot.S_LOCATION, "PDF → pregunta ubicación")
    await drive(upd_cq(f"loc_{bot.LOCATIONS[1]}", fake.last_message_id))
    await drive(upd_cq(f"cat_{bot.CATEGORIES[1]}", fake.last_message_id))
    items = ud.get("batch_items") or []
    check(len(items) == 2,
          f"2 páginas del PDF procesadas como facturas (hay {len(items)})")
    check(state() == bot.S_BATCH_LIST, "lista del lote PDF (S_BATCH_LIST)")
    await drive(upd_cq("baccall", fake.last_message_id))
    guardadas = bot.get_facturas("2026-06")
    check(len(guardadas) == 9,
          f"total 9 facturas guardadas (hay {len(guardadas)})")

    print("\n— Escenario F: fechas en formato latino (DD/MM/AAAA) —")
    nf = bot.normalize_fecha
    check(nf("15/06/2026") == "2026-06-15", "acepta DD/MM/AAAA")
    check(nf("15-06-2026") == "2026-06-15", "acepta DD-MM-AAAA")
    check(nf("5/6/26") == "2026-06-05", "acepta D/M/AA")
    check(nf("2026-06-15") == "2026-06-15", "ISO pasa igual")
    check(nf("31/02/x") == "" and nf("ayer") == "", "basura no pasa")
    check(bot.fecha_display("2026-06-15") == "15/06/2026",
          "se muestra como 15/06/2026")
    vf = bot.validate_and_fix({"_error": None, "fecha_comprobante": "15/06/2026",
                               "total_facturado": 118, "itbis": 18,
                               "monto_sin_itbis": 100, "rnc": "101000001",
                               "ncf": "B0100000099"})
    check(vf["fecha_comprobante"] == "2026-06-15",
          "validate_and_fix normaliza fecha latina de la IA")

    # Flujo real: corregir la fecha escribiéndola en formato latino
    await drive(upd_photo("FCHA1"))
    await drive(upd_cq(f"loc_{bot.LOCATIONS[0]}", fake.last_message_id))
    await drive(upd_cq(f"cat_{bot.CATEGORIES[0]}", fake.last_message_id))
    await drive(upd_cq("confirm_edit", fake.last_message_id))
    await drive(upd_cq("edit_fecha", fake.last_message_id))
    await drive(upd_text("no-es-fecha"))
    check(state() == bot.S_EDIT_VALUE,
          "fecha inválida se rechaza y sigue pidiendo el valor")
    await drive(upd_text("20/06/2026"))
    data = ud.get("pending_invoice") or {}
    check(data.get("fecha_comprobante") == "2026-06-20",
          "fecha editada '20/06/2026' queda ISO internamente")
    check(state() == bot.S_CONFIRM, "vuelve a la tarjeta de revisión")
    card_params = [p for a, p in fake.calls
                   if a == "sendMessage" and "Fecha" in p.get("text", "")][-1]
    check("20/06/2026" in card_params["text"],
          "la tarjeta muestra la fecha en formato latino")
    await drive(upd_cq("confirm_accept", fake.last_message_id))
    check(len(bot.get_facturas("2026-06")) == 10,
          "quedó guardada en el mes correcto (2026-06)")
    await drive(upd_cq("fin_lote", fake.last_message_id))

    print("\n— Escenario C: escape de Markdown —")
    nombre_feo = "FERRE_MAX *SRL* [STO DGO]"
    check(hasattr(bot, "md"), "existe helper bot.md() de escape")
    if hasattr(bot, "md"):
        esc = bot.md(nombre_feo)
        check("\\_" in esc and "\\*" in esc and "\\[" in esc,
              "md() escapa _ * [")
        card = bot.format_review_message(
            {"nombre_proveedor": nombre_feo, "ncf": "B0100000001",
             "rnc": "101000001", "fecha_comprobante": "2026-06-15",
             "monto_sin_itbis": 100, "itbis": 18, "propina": 0,
             "total_facturado": 118, "metodo_pago": "EFECTIVO",
             "nivel_confianza": "ALTO"},
            "Punta Cana", "Casa", "2026-06", "02")
        check("FERRE\\_MAX" in card,
              "la tarjeta de revisión escapa el nombre del proveedor")

    # ── Escenario G: validaciones que atajan lo que la DGII rechaza ──────
    print("\n— Escenario G: RNC, fechas imposibles y NCF repetido —")

    # Dígito verificador del RNC. El caso real: '102060621' se coló como
    # Bellón SAS en el 606 de julio y la DGII rechazó la línea; el bueno
    # es '102000621'. Los dos tienen 9 dígitos, así que el largo no basta.
    check(bot.rnc_valido("102000621"), "RNC bueno de Bellón pasa")
    check(not bot.rnc_valido("102060621"), "RNC malo de Bellón NO pasa")
    for r in ("131545157", "133311887", "131092659", "132859146"):
        check(bot.rnc_valido(r), f"RNC real {r} pasa")
    check(bot.rnc_valido("40212345678"), "cédula de 11 dígitos se da por buena")
    check(not bot.rnc_valido("12345"), "RNC corto no pasa")

    d = bot.validate_and_fix({"rnc": "102060621", "ncf": "E310001373766",
                              "total_facturado": 616.86, "itbis": 94.10,
                              "monto_sin_itbis": 522.76, "propina": 0})
    check(any("dígito verificador" in w for w in d["_warnings"]),
          "validate_and_fix advierte del RNC inválido")
    check(d["_needs_review"], "y la deja marcada para revisar")

    # Fechas que no existen en el calendario
    check(bot.normalize_fecha("31/02/2026") == "", "31/02 se rechaza")
    check(bot.normalize_fecha("31/04/2026") == "", "31/04 se rechaza")
    check(bot.normalize_fecha("29/02/2026") == "", "29/02 en año no bisiesto se rechaza")
    check(bot.normalize_fecha("29/02/2024") == "2024-02-29", "29/02 bisiesto sí pasa")
    check(bot.normalize_fecha("15/07/2026") == "2026-07-15", "fecha normal pasa")

    # El NCF repetido lo para la BASE, no solo el chequeo de aplicación
    fac = {"rnc": "131545157", "ncf": "E310009999999", "nombre_proveedor": "PRUEBA",
           "total_facturado": 100, "itbis": 0, "monto_sin_itbis": 100, "propina": 0}
    id1 = bot.save_factura("2026-06", "Punta Cana", "Obra", dict(fac))
    id2 = bot.save_factura("2026-06", "Punta Cana", "Obra", dict(fac))
    check(id1 > 0, "la primera se guarda")
    check(id2 == 0, "la segunda la rechaza el índice único (devuelve 0)")

    # /pendientes: la advertencia se conserva; lo que la saca de la cola es
    # que un humano la haya revisado (revisada_manual), no haberla aceptado.
    con_adv = dict(fac, ncf="E310009999998", _needs_review=True,
                   _warnings=["⚠ prueba"])
    bot.save_factura("2026-06", "Punta Cana", "Obra", dict(con_adv), reviewed=False)
    bot.save_factura("2026-06", "Punta Cana", "Obra",
                     dict(con_adv, ncf="E310009999997"), reviewed=True)
    fs = {f["ncf"]: f for f in bot.get_facturas("2026-06")}
    check(fs["E310009999998"]["needs_review"] == 1,
          "sin revisar: queda como pendiente")
    check(fs["E310009999997"]["needs_review"] == 1
          and fs["E310009999997"]["revisada_manual"] == 1,
          "revisada: conserva la advertencia pero sale de la cola")

    # ── Escenario H: el grupo espejo ────────────────────────────────────
    print("\n— Escenario H: aviso al grupo espejo —")
    GRUPO = "-1009999999999"
    factura = dict(nombre_proveedor="Bellón S.A.S", ncf="E310001364339",
                   rnc="102000621", total_facturado=10500.0, itbis=1601.69)

    def avisos_al_grupo():
        return [p for m, p in fake.calls
                if m == "sendMessage" and str(p.get("chat_id")) == GRUPO]

    n0 = len(avisos_al_grupo())
    bot.GROUP_CHAT_ID = ""
    await bot.notify_group(app, "no debería salir")
    check(len(avisos_al_grupo()) == n0, "sin GROUP_CHAT_ID no se avisa a nadie")

    bot.GROUP_CHAT_ID = GRUPO

    class _Ctx:
        bot = tbot
    await bot.notify_group(_Ctx, "hola grupo")
    check(len(avisos_al_grupo()) == n0 + 1, "con GROUP_CHAT_ID sí se avisa")

    await bot.notify_group(_Ctx, "eco", origen_chat_id=GRUPO)
    check(len(avisos_al_grupo()) == n0 + 1,
          "un mensaje que nace en el grupo no se reenvía al grupo (sin eco)")

    texto = bot._aviso_factura(7, factura, "2026-08", "02", "@Cormurca")
    check("#7" in texto and "@Cormurca" in texto and "E310001364339" in texto,
          "el aviso trae id, quién la subió y el NCF")
    check("10,500.00" in texto, "y el monto formateado")

    con_adv = dict(factura, ncf="E310001364340",
                   _warnings=["RNC no pasa el dígito verificador"])
    check("⚠️" in bot._aviso_factura(8, con_adv, "2026-08", "02", "@ricfut"),
          "las advertencias se ven en el grupo")

    items = [{"status": "accepted", "data": dict(factura, ncf=f"E31000000{i:04d}",
                                                 total_facturado=100.0)}
             for i in range(30)]
    items.append({"status": "discarded", "data": dict(factura)})
    lote = bot._aviso_lote(items, {"2026-08"}, "@Cormurca", dups=2)
    check("30 factura(s)" in lote, "el aviso de lote cuenta solo las guardadas")
    check("…y 5 más" in lote, "y corta la lista en 25 para no pasar el tope de Telegram")
    check("3,000.00" in lote, "con el total del lote")
    check("2 duplicada(s)" in lote, "y las duplicadas omitidas")
    check(bot._aviso_lote([], set(), "@x", 0) == "",
          "un lote sin nada guardado no avisa")

    # Que un fallo de Telegram no tumbe el guardado
    class _CtxRoto:
        class bot:
            @staticmethod
            async def send_message(**kw):
                raise RuntimeError("el bot no está en el grupo")
    await bot.notify_group(_CtxRoto, "esto explota por dentro")
    check(True, "si el envío al grupo falla, no se propaga la excepción")

    bot.GROUP_CHAT_ID = ""

    # ── Escenario I: el Excel de comprobantes del Banco Santa Cruz ───────
    print("\n— Escenario I: reporte del Banco Santa Cruz —")
    from openpyxl import Workbook
    from datetime import datetime

    def bsc_xlsx(filas, encabezados=("NCF", "Moneda", "Monto", "Fecha Generación"),
                 fila_hdr=2):
        wb = Workbook(); ws = wb.active; ws.title = "NcfSummary-export-20260917"
        for c, h in enumerate(encabezados, 1):
            ws.cell(fila_hdr, c, h)
        for i, f in enumerate(filas):
            for c, v in enumerate(f, 1):
                ws.cell(fila_hdr + 1 + i, c, v)
        b = io.BytesIO(); wb.save(b); return b.getvalue()

    d = datetime(2026, 8, 3)
    filas, probs = bot.parse_bsc_excel(bsc_xlsx([
        ("E310004470672", "RD$", 563.8, d),
        ("E310004469302", "RD$", 200,   datetime(2026, 8, 31)),
    ]))
    check(len(filas) == 2 and not probs, "lee el formato del banco tal cual")
    r = filas[0]
    check(r["rnc"] == "102012921" and r["nombre_proveedor"] == bot.NOMBRE_BSC,
          "les pone el RNC y el nombre del banco")
    check(r["itbis"] == 0 and r["monto_sin_itbis"] == 563.8
          and r["total_facturado"] == 563.8,
          "sin ITBIS: el monto completo es la base")
    check(r["fecha_comprobante"] == "2026-08-03", "la fecha sale en ISO")
    check(filas[1]["fecha_comprobante"][:7] == "2026-08", "y el mes sale de la fecha")

    # encabezados movidos de sitio y en otro orden
    filas2, _ = bot.parse_bsc_excel(bsc_xlsx(
        [(d, 563.8, "RD$", "E310004470672")],
        encabezados=("Fecha Generacion", "Monto", "Moneda", "NCF"), fila_hdr=5))
    check(len(filas2) == 1 and filas2[0]["ncf"] == "E310004470672",
          "encuentra las columnas por nombre aunque el banco las mueva")

    filas3, probs3 = bot.parse_bsc_excel(bsc_xlsx([
        ("E310004470672", "RD$", 563.8, d),
        ("E310004469999", "US$", 100.0, d),      # otra moneda
        ("E310004469998", "RD$", None,  d),      # sin monto
        ("E310004469997", "RD$", -5.0,  d),      # monto negativo
        ("NO-ES-UN-NCF",  "RD$", 50.0,  d),      # NCF con mala forma
        ("E310004469996", "RD$", 50.0,  "nunca"),  # fecha ilegible
    ]))
    check(len(filas3) == 1, "descarta dólares, montos malos, NCF raros y fechas ilegibles")
    check(len(probs3) == 5, "y explica cada fila que omitió")
    check(any("US$" in p for p in probs3), "diciendo cuál venía en otra moneda")

    vacio, probs4 = bot.parse_bsc_excel(bsc_xlsx([], encabezados=("A", "B", "C")))
    check(not vacio and probs4, "un archivo que no es del banco se rechaza con explicación")
    check(bot.parse_bsc_excel(b"no soy un xlsx")[0] == [],
          "y un archivo corrupto no revienta")

    # guardar de verdad, con duplicados
    guardadas = 0
    for r in filas:
        if bot.save_factura(r["fecha_comprobante"][:7], "Santo Domingo", "Banco",
                            r, "@ricfut", tipo_gasto="07", reviewed=True):
            guardadas += 1
    check(guardadas == 2, "las dos se guardan")
    check(bot.save_factura("2026-08", "Santo Domingo", "Banco", filas[0],
                           "@ricfut", tipo_gasto="07", reviewed=True) == 0,
          "y reenviar el mismo reporte no las duplica")
    fs = {f["ncf"]: f for f in bot.get_facturas("2026-08")}
    check(fs["E310004470672"]["tipo_gasto"] == "07",
          "quedan como 07 — gastos financieros")
    check(fs["E310004470672"]["needs_review"] == 0,
          "y no caen en la cola de pendientes")

    # ── Escenario J: /id y el silencio del bot en grupos ────────────────
    print("\n— Escenario J: /id y el bot en un grupo —")
    GID = -1001234567890

    def upd_grupo(text, uid=UID):
        return {"update_id": next(_upd_id),
                "message": {"message_id": next(_msg_id), "date": int(time.time()),
                            "chat": {"id": GID, "type": "supergroup",
                                     "title": "606 ROMUR"},
                            "from": {"id": uid, "is_bot": False,
                                     "first_name": "Ricardo"},
                            "text": text,
                            "entities": [{"type": "bot_command",
                                          "offset": 0, "length": len(text)}]
                            if text.startswith("/") else []}}

    n0 = len(fake.calls)
    await drive(upd_grupo("/id"))
    envios = [p for m, p in fake.calls[n0:] if m == "sendMessage"]
    check(len(envios) == 1 and str(GID) in envios[0]["text"],
          "/id en el grupo devuelve el ID del grupo")
    check("GROUP_CHAT_ID" in envios[0]["text"],
          "y te dice cómo ponerlo en Railway")

    # un usuario sin desbloquear no debe provocar que el bot pida la
    # contraseña delante de todo el grupo
    OTRO = 999
    bot.ALLOWED_USERS = set()
    n1 = len(fake.calls)
    await drive(upd_grupo("hola gente", uid=OTRO))
    check(not [p for m, p in fake.calls[n1:] if m == "sendMessage"],
          "el bot no pide la contraseña ni contesta nada en el grupo")
    check(app.user_data[OTRO].get("unlocked") is not True,
          "y tampoco desbloquea a nadie por escribir en el grupo")

    n2 = len(fake.calls)
    await drive(upd_grupo("/id", uid=OTRO))
    check([p for m, p in fake.calls[n2:] if m == "sendMessage"],
          "pero /id sí pasa aunque no haya desbloqueado (hace falta para configurar)")

    # ── Escenario K: el archivo de envío de la DGII ─────────────────────
    print("\n— Escenario K: archivo de envío 606 —")

    def fac(**kw):
        base = dict(rnc="102000621", ncf="E310001364339", nombre="Bellón",
                    fecha_comp="2026-08-14", base=1000.0, itbis=180.0,
                    propina=0.0, total=1180.0, tipo_gasto="02",
                    metodo="TARJETA_CREDITO")
        base.update(kw); base["total"] = base["base"] + base["itbis"] + base["propina"]
        return base

    filas, bloq, avi = bot.preparar_606([fac()], "2026-08")
    check(not bloq and not avi and len(filas) == 1, "una factura limpia pasa sin ruido")
    linea = bot.construir_txt_606(filas, "131545157", "202608").decode().splitlines()
    check(linea[0] == "606|131545157|202608|1", "la cabecera trae RNC, período y conteo")
    check(linea[1] == "102000621|1|02|E310001364339||20260814|20260814|"
                      "1000||1000|180|||180|0|||||||0|03",
          "y la línea sale con el formato exacto de la Herramienta DGII")

    # el tope de la propina: 10% truncado, nunca redondeado
    f, _, a = bot.preparar_606([fac(base=1292.37, itbis=232.63, propina=129.24,
                                    tipo_gasto="05")], "2026-08")
    check(f[0]["propina"] == 129.23, "la propina se trunca al 10% (129.24 → 129.23)")
    check(any("pasa del 10%" in x for x in a), "y avisa de que la bajó")
    check(bot.tope_propina(1292.37) == 129.23, "tope_propina trunca, no redondea")

    # propina ⇒ tipo 05
    f, _, a = bot.preparar_606([fac(propina=50.0, base=500.0, itbis=90.0)], "2026-08")
    check(f[0]["tipo_gasto"] == "05", "con propina, el tipo pasa a 05")

    # 07 solo para el Banco Santa Cruz
    f, _, a = bot.preparar_606([fac(tipo_gasto="07")], "2026-08")
    check(f[0]["tipo_gasto"] == "02", "un 07 que no es del banco se reclasifica a 02")
    f, _, _ = bot.preparar_606([fac(rnc=bot.RNC_BSC, ncf="E310004470672",
                                    tipo_gasto="07", itbis=0.0,
                                    metodo="NOTA_CREDITO")], "2026-08")
    check(f[0]["forma_pago"] == "06", "el del banco va con forma de pago 06")

    # los financieros al final
    f, _, _ = bot.preparar_606([
        fac(rnc=bot.RNC_BSC, ncf="E310004470672", tipo_gasto="07", itbis=0.0),
        fac(ncf="E310001364340"),
    ], "2026-08")
    check([x["tipo_gasto"] for x in f] == ["02", "07"],
          "los gastos financieros se van al final del archivo")

    # bloqueos: lo que la DGII rechaza seguro
    _, bloq, _ = bot.preparar_606([fac(rnc="101013814")], "2026-08")
    check(any("no existe" in b for b in bloq), "un RNC con dígito verificador malo bloquea")
    _, bloq, _ = bot.preparar_606([fac(ncf="E31000066076")], "2026-08")
    check(any("mal formado" in b for b in bloq), "un NCF con largo raro bloquea")
    _, bloq, _ = bot.preparar_606([fac(), fac()], "2026-08")
    check(any("repetido" in b for b in bloq), "un NCF repetido bloquea")
    _, bloq, _ = bot.preparar_606([fac(fecha_comp="")], "2026-08")
    check(any("sin fecha" in b for b in bloq), "sin fecha de comprobante, bloquea")

    # comprobante de otro mes: avisa pero no bloquea
    _, bloq, a = bot.preparar_606([fac(fecha_comp="2026-07-30")], "2026-08")
    check(not bloq and any("no de 202608" in x for x in a),
          "un comprobante de otro mes avisa pero deja pasar")

    # formas de pago
    check(bot.forma_pago_dgii("EFECTIVO", "02", "102000621") == "01", "efectivo → 01")
    check(bot.forma_pago_dgii("CHEQUE", "02", "102000621") == "02", "cheque → 02")
    check(bot.forma_pago_dgii("TRANSFERENCIA", "02", "102000621") == "02",
          "transferencia → 02")
    check(bot.forma_pago_dgii("TARJETA_DEBITO", "02", "102000621") == "03",
          "débito → 03 (no 05: los códigos de Gabi no son los de la DGII)")
    check(bot.forma_pago_dgii("", "02", "102000621") == "03",
          "sin método conocido, el default es tarjeta")

    check(bot.fmt_txt_num(724100.0) == "724100" and bot.fmt_txt_num(13033.80) == "13033.8"
          and bot.fmt_txt_num(0) == "0" and bot.fmt_txt_num(169.49) == "169.49",
          "los montos van sin ceros de cola, como los escribe la Herramienta")

    await app.shutdown()

    print()
    if FAILS:
        print(f"❌ {len(FAILS)} verificación(es) fallaron:")
        for f in FAILS:
            print(f"   - {f}")
        return 1
    print("✅ Todas las verificaciones pasaron.")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(run()))
